"""PDF and DOCX text extraction and chunking for knowledge-graph generation."""

from __future__ import annotations

import io
import math
import os
from typing import Any, Dict, List, Literal, Tuple

import fitz  # PyMuPDF
from docx import Document
from fastapi import HTTPException, status

SupportedKind = Literal["pdf", "docx"]

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx"}


def get_max_pages() -> int:
    return int(os.getenv("KG_MAX_PAGES", "20"))


def get_max_file_mb() -> int:
    return int(os.getenv("KG_MAX_FILE_MB", "20"))


def get_chunk_page_size() -> int:
    return int(os.getenv("KG_CHUNK_PAGES", "5"))


def get_docx_words_per_page() -> int:
    return max(1, int(os.getenv("KG_DOCX_WORDS_PER_PAGE", "500")))


def detect_file_kind(filename: str | None) -> SupportedKind:
    name = (filename or "").lower()
    for ext, kind in SUPPORTED_EXTENSIONS.items():
        if name.endswith(ext):
            return kind  # type: ignore[return-value]
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Only PDF and DOCX files are supported",
    )


def validate_upload(filename: str | None, content: bytes) -> SupportedKind:
    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty",
        )
    kind = detect_file_kind(filename)
    max_bytes = get_max_file_mb() * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File exceeds maximum size of {get_max_file_mb()} MB",
        )
    return kind


def _chunk_page_texts(page_texts: List[str]) -> List[Dict[str, Any]]:
    page_count = len(page_texts)
    chunk_size = max(1, get_chunk_page_size())
    chunks: List[Dict[str, Any]] = []
    for start_idx in range(0, page_count, chunk_size):
        end_idx = min(start_idx + chunk_size, page_count)
        text = "\n\n".join(page_texts[start_idx:end_idx]).strip()
        chunks.append(
            {
                "page_start": start_idx + 1,
                "page_end": end_idx,
                "text": text,
            }
        )
    return chunks


def _enforce_page_limit(page_count: int, label: str) -> None:
    if page_count < 1:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"{label} has no extractable text",
        )
    max_pages = get_max_pages()
    if page_count > max_pages:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"{label} has {page_count} pages; maximum allowed is {max_pages}",
        )


def extract_pdf_pages_and_chunks(content: bytes) -> Tuple[int, List[Dict[str, Any]]]:
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid or unreadable PDF: {exc}",
        ) from exc

    try:
        page_count = doc.page_count
        _enforce_page_limit(page_count, "PDF")
        page_texts = [
            (doc.load_page(i).get_text("text") or "") for i in range(page_count)
        ]
        return page_count, _chunk_page_texts(page_texts)
    finally:
        doc.close()


def _docx_paragraph_and_table_text(doc: Document) -> str:
    parts: List[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts).strip()


def _split_text_into_estimated_pages(full_text: str) -> List[str]:
    words = full_text.split()
    if not words:
        return []
    per_page = get_docx_words_per_page()
    pages: List[str] = []
    for start in range(0, len(words), per_page):
        pages.append(" ".join(words[start : start + per_page]))
    return pages


def extract_docx_pages_and_chunks(content: bytes) -> Tuple[int, List[Dict[str, Any]]]:
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid or unreadable DOCX: {exc}",
        ) from exc

    full_text = _docx_paragraph_and_table_text(doc)
    if not full_text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="DOCX has no extractable text",
        )

    page_texts = _split_text_into_estimated_pages(full_text)
    page_count = len(page_texts)
    # Prefer ceil(words/per_page) for messaging consistency
    word_count = len(full_text.split())
    estimated = max(1, math.ceil(word_count / get_docx_words_per_page()))
    _enforce_page_limit(estimated, "DOCX (estimated)")
    return page_count, _chunk_page_texts(page_texts)


def extract_pages_and_chunks(
    content: bytes,
    kind: SupportedKind,
) -> Tuple[int, List[Dict[str, Any]]]:
    """
    Returns (page_count, chunks) where each chunk is:
    { page_start, page_end, text }

    For DOCX, page_count is estimated from word count (~KG_DOCX_WORDS_PER_PAGE words/page).
    """
    if kind == "pdf":
        return extract_pdf_pages_and_chunks(content)
    return extract_docx_pages_and_chunks(content)


def file_extension(kind: SupportedKind) -> str:
    return ".pdf" if kind == "pdf" else ".docx"
